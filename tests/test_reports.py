import datetime as dt

import openpyxl
import pytest

from aoiro_fy12_to_fy3.cli import main
from aoiro_fy12_to_fy3.ledger import compute_terms
from aoiro_fy12_to_fy3.report_financials import software_values, write_official, write_software
from aoiro_fy12_to_fy3.verify import normalize, parse_software_financials_xls
from aoiro_fy12_to_fy3.report_ledger_pdf import FontNotFoundError, account_rows, find_font, write_ledger_pdf
from aoiro_fy12_to_fy3.statements import build_financials


def _labels(ws) -> dict[str, list]:
    """シートの「項目名 → その行の数値」。"""
    out = {}
    for row in ws.iter_rows():
        label = None
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip():
                label = cell.value.strip()
            elif isinstance(cell.value, int) and label:
                out.setdefault(label, []).append(cell.value)
                label = None
    return out


def test_financials_identities(config):
    for r in compute_terms(config, 2018):
        fin = build_financials(r)
        assert fin.asset_total == fin.liabilities_and_capital_total
        assert fin.net_assets_subtotal == fin.retained_closing + fin.reserve_total
        assert fin.retained_opening + fin.reserve_total + fin.reserve_transfer_total + fin.net_income == (
            fin.capital_total
        )


def test_official_xlsx(config, tmp_path):
    (r,) = compute_terms(config, 2017)
    fin = build_financials(r)
    path = tmp_path / "official.xlsx"
    write_official(path, config, fin)
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["表紙", "貸借対照表", "損益計算書"]
    cover = [c.value for row in wb["表紙"].iter_rows() for c in row if c.value]
    assert "第 22期" in cover and "自 2017年 4月 1日" in cover and "至 2018年 3月31日" in cover
    assert "架空会" in cover

    bs = _labels(wb["貸借対照表"])
    # 正式様式の「繰越正味財産」は期首の額（期中の積立金振替は別の行に出す）
    assert bs["繰越正味財産"] == [510000]
    assert bs["記念行事積立金"] == [60000]
    assert bs["建替積立金"] == [60000]  # label で表示名を変えられる
    assert bs["積立金　計"] == [120000]
    assert bs["当期積立金振替額"] == [-20000]
    assert bs["負債・正味財産の部合計"] == [fin.asset_total]

    pl = _labels(wb["損益計算書"])
    assert pl["会費収入"] == [40000]
    assert pl["その他収入"] == [5]
    assert pl["援助費（OB・OG）"] == [0]  # always_show
    assert "通信費" not in pl  # 0円の行は除外済み
    assert pl["一般管理費  計"] == [-3432]
    assert pl["当期正味財産増減額"] == [fin.net_income]
    assert pl["当期正味財産増減額－当期積立金振替額"] == [fin.net_income - 20000]


def _xls_labels(sheet) -> dict[str, list]:
    out = {}
    for r in range(sheet.nrows):
        label = None
        for v in sheet.row_values(r):
            if isinstance(v, str) and v.strip():
                label = v.strip()
            elif isinstance(v, float) and label:
                out.setdefault(label, []).append(int(v))
                label = None
    return out


def test_software_xls(config, tmp_path):
    import xlrd

    (r,) = compute_terms(config, 2017)
    fin = build_financials(r)
    path = tmp_path / "software.xls"
    write_software(path, config, fin)
    wb = xlrd.open_workbook(str(path))
    assert wb.sheet_names() == ["表紙", "貸借対照表", "損益計算書", "販売費及び一般管理費内訳書"]
    bs = _xls_labels(wb.sheet_by_name("貸借対照表"))
    assert bs["繰越正味財産"] == [490000]
    assert bs["正味財産  計"] == [610000]
    assert bs["資本の部合計"] == [610000 + fin.net_income]
    sga = _xls_labels(wb.sheet_by_name("販売費及び一般管理費内訳書"))
    assert sga["支払手数料"] == [432]
    assert sga["販売費・一般管理費計"] == [3432]
    values = software_values(fin)
    assert values["任意積立金合計"] == 120000
    assert values["売上高計"] == 40005
    # ソフト出力の決算書xlsと同じ読み方で、ツールの数値がそのまま読み取れる
    parsed = parse_software_financials_xls(path)
    assert {normalize(k): v for k, v in values.items()}.items() <= parsed.items()


def test_ledger_rows(config):
    (r,) = compute_terms(config, 2017)
    rows = account_rows(r, 106)
    assert rows[0].kind == "opening" and rows[0].balance == 700000
    subtotals = [row.memo for row in rows if row.kind == "subtotal"]
    assert subtotals == ["5月度合計", "6月度合計", "11月度合計", "12月度合計", "2月度合計"]
    total = rows[-1]
    assert total.kind == "total" and total.sub_memo == "6件"
    entries = [row for row in rows if row.kind == "entry"]
    assert entries[-1].balance == r.closing_balance(106)
    # 複合仕訳の相手科目は「複合」
    assert any(row.counter == "複合" for row in entries)
    # 同じ日付が続くときは最後の行にだけ残高を出す
    may = [row for row in entries if row.date == dt.date(2017, 5, 10)]
    assert [row.show_balance for row in may] == [False, True]

    closing = account_rows(r, 409)
    assert [row.memo for row in closing if row.kind == "subtotal"] == ["決算整理合計"]


def _font_available() -> bool:
    try:
        find_font()
    except FontNotFoundError:
        return False
    return True


@pytest.mark.skipif(not _font_available(), reason="日本語フォントがない")
def test_ledger_pdf(config, tmp_path):
    pdfplumber = pytest.importorskip("pdfplumber")
    (r,) = compute_terms(config, 2017)
    path = tmp_path / "ledger.pdf"
    pages = write_ledger_pdf(path, config.name, r, created=dt.date(2026, 1, 1))
    assert pages == len(r.accounts_in_use())
    with pdfplumber.open(path) as pdf:
        text = "\n".join(p.extract_text() for p in pdf.pages)
    assert "106:普通預金" in text
    assert "(2017年4月1日～2018年3月31日)" in text
    assert "総 合 計" in text or "総　合　計" in text
    assert "郵送料（訂正）" not in text  # 0円の行


@pytest.mark.skipif(not _font_available(), reason="日本語フォントがない")
def test_cli_build(org_dir, tmp_path, capsys):
    out = tmp_path / "out"
    assert main(["build", "--config", str(org_dir / "org.yaml"), "--all", "--out", str(out)]) == 0
    names = sorted(p.relative_to(out).as_posix() for p in out.rglob("*") if p.is_file())
    assert "総勘定元帳/総勘定元帳(架空会_第23期).pdf" in names
    assert "決算書/決算書(架空会_第22期).xlsx" in names
    assert "決算書/決算書(架空会_第22期).xls" in names
    assert "金額0円の行を元帳から除外: 伝票3" in capsys.readouterr().err
