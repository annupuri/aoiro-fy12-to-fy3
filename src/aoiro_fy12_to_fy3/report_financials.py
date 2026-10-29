"""決算書の出力。

2つの様式を出す。
- official: 総会で報告してきた過去の正式な決算書（第20・21期）の様式。xlsx
- software: 『みんなの青色申告』が出力する決算書の様式。ソフトの出力と突き合わせる検証用。xls

シートの内容はいったん Sheet に組み立て、xlsx（openpyxl）または xls（xlwt）で書き出す。
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from .config import OrgConfig
from .statements import Financials

MINCHO = "ＭＳ 明朝"
GOTHIC = "ＭＳ ゴシック"
NUMBER_FORMAT = '#,##0;"△ "#,##0'
TAX_NOTE = "[税込]（単位：円）"
BS_SHEET = "貸借対照表"


@dataclass
class Cell:
    value: object = None
    font: str = MINCHO
    size: float = 11
    bold: bool = False
    align: str | None = None
    number: bool = False
    borders: set[str] = field(default_factory=set)


@dataclass
class Sheet:
    """書式付きのセルの集まり。行・列は1始まり。"""

    title: str
    cells: dict[tuple[int, int], Cell] = field(default_factory=dict)
    merges: list[tuple[int, int, int, int]] = field(default_factory=list)
    widths: dict[int, float] = field(default_factory=dict)

    def cell(self, row: int, col: int) -> Cell:
        return self.cells.setdefault((row, col), Cell())

    def put(self, row: int, col: int, value, *, bold=False, align=None, number=False, size=11, font=MINCHO):
        c = self.cell(row, col)
        c.value, c.bold, c.align, c.number, c.size, c.font = value, bold, align, number, size, font
        return c

    def merge(self, row1: int, col1: int, row2: int, col2: int) -> None:
        self.merges.append((row1, col1, row2, col2))

    def box(self, top: int, bottom: int, left: int, right: int) -> None:
        """範囲の外枠と列の区切り線を引く。"""
        for r in range(top, bottom + 1):
            for c in range(left, right + 1):
                sides = self.cell(r, c).borders
                sides |= {"left", "right"}
                if r == top:
                    sides.add("top")
                if r == bottom:
                    sides.add("bottom")

    def underline(self, row: int, left: int, right: int) -> None:
        for c in range(left, right + 1):
            self.cell(row, c).borders.add("bottom")


@dataclass(frozen=True)
class Line:
    """表の1行。value_col は金額を入れる列（ラベル列の右隣からの位置）。value が None なら見出し行。"""

    label: str
    value: int | None = None
    value_col: int = 0
    bold: bool = False


def _date_jp(year: int, month: int, day: int) -> str:
    return f"{year}年{month:2d}月{day:2d}日"


def period_text(fin: Financials) -> str:
    t = fin.result.term
    return f"自 {_date_jp(t.start.year, 4, 1)}  至 {_date_jp(t.end.year, 3, 31)}"


def as_of_text(fin: Financials) -> str:
    return f"{_date_jp(fin.result.term.end.year, 3, 31)} 現在"


def _title(ws: Sheet, text: str, last_col: int, org: str, period: str) -> None:
    ws.merge(1, 2, 1, last_col)
    ws.put(1, 2, text, bold=True, size=16, font=GOTHIC, align="center")
    ws.put(2, last_col, TAX_NOTE, size=9, font=GOTHIC, align="right")
    ws.put(3, 2, org)
    ws.put(3, last_col, period, align="right")
    ws.underline(3, 2, last_col)


def _widths(ws: Sheet, widths: dict[int, float]) -> None:
    ws.widths = {1: 2.9, **widths}


def _lines(ws: Sheet, lines: list[Line], start: int, label_col: int = 2) -> int:
    """1列のラベルと金額の表を書き、最後の行番号を返す。"""
    for i, line in enumerate(lines):
        ws.put(start + i, label_col, line.label, bold=line.bold, align="center" if line.bold else "left")
        if line.value is not None:
            ws.put(start + i, label_col + 1 + line.value_col, line.value, bold=line.bold, number=True)
    return start + len(lines) - 1


def cover_sheet(config: OrgConfig, fin: Financials) -> Sheet:
    ws = Sheet("表紙")
    _widths(ws, {2: 9.2})
    t = fin.result.term

    def centered(row: int, text: str, left: int, right: int, **style) -> None:
        ws.merge(row, left, row, right)
        ws.put(row, left, text, align="center", **style)

    centered(13, "決　算　報　告　書", 2, 10, size=28, font=GOTHIC)
    centered(19, f"第 {t.number}期", 5, 7, bold=True, size=14)
    centered(23, f"自 {_date_jp(t.start.year, 4, 1)}", 5, 7)
    centered(25, f"至 {_date_jp(t.end.year, 3, 31)}", 5, 7)
    centered(40, config.name, 2, 10)
    for i, line in enumerate(config.address):
        centered(44 + i, line, 2, 10)
    return ws


def _balance_sheet(config: OrgConfig, fin: Financials, right_title: str, left: list[Line], right: list[Line],
                   left_total: Line, right_total: Line) -> Sheet:
    """勘定式の貸借対照表。左右の表の高さをそろえ、最下行に合計を置く。"""
    ws = Sheet(BS_SHEET)
    _widths(ws, {2: 30.6, 3: 17.5, 4: 30.6, 5: 17.5})
    _title(ws, "貸　借　対　照　表", 5, config.name, as_of_text(fin))
    ws.merge(4, 2, 4, 3)
    ws.merge(4, 4, 4, 5)
    ws.put(4, 2, "資　産　の　部", bold=True, align="center")
    ws.put(4, 4, right_title, bold=True, align="center")
    for col in (2, 4):
        ws.put(5, col, "科　　目", align="center")
        ws.put(5, col + 1, "金　額", align="center")
    for left_col in (2, 4):
        ws.box(4, 5, left_col, left_col + 1)
        ws.box(5, 5, left_col, left_col + 1)

    n = max(len(left), len(right))
    left = left + [Line("")] * (n - len(left)) + [left_total]
    right = right + [Line("")] * (n - len(right)) + [right_total]
    start = 6
    end = _lines(ws, left, start, 2)
    _lines(ws, right, start, 4)
    ws.box(start, end, 2, 3)
    ws.box(start, end, 4, 5)
    ws.underline(end - 1, 2, 5)
    return ws


# ---------------------------------------------------------------- 正式様式


def official_balance_sheet(config: OrgConfig, fin: Financials) -> Sheet:
    left = [Line("  【流動資産】"), Line("    （現金・預金）")]
    left += [Line(f"      {a.name}", v) for a, v in fin.assets]
    left += [Line("        現金・預金 計", fin.asset_total), Line("          流動資産合計", fin.asset_total)]

    right = [Line("  【流動負債】")]
    right += [Line(f"      {a.name}", v) for a, v in fin.liabilities]
    right += [
        Line("          流動負債  計", fin.liability_total),
        Line("負債の部合計", fin.liability_total, bold=True),
        Line("正　味　財　産　の　部", bold=True),
        Line("  【正味財産】"),
        Line(f"    {fin.result.chart.retained.name}", fin.retained_opening),
        Line("    積立金"),
    ]
    right += [Line(f"     {label}", v) for label, v in fin.reserve_rows]
    right += [
        Line("      積立金　計", fin.reserve_total),
        Line("    　当期積立金振替額", fin.reserve_transfer_total),
        Line("    当期正味財産増減額", fin.net_income),
        Line("正味財産の部合計", fin.capital_total, bold=True),
    ]
    return _balance_sheet(
        config, fin, "負　債　・　正　味　財　産　の　部", left, right,
        Line("資産の部合計", fin.asset_total, bold=True),
        Line("負債・正味財産の部合計", fin.liabilities_and_capital_total, bold=True),
    )


def official_income_statement(config: OrgConfig, fin: Financials) -> Sheet:
    ws = Sheet("損益計算書")
    _widths(ws, {2: 47.1, 3: 20.0, 4: 20.0})
    _title(ws, "損　益　計　算　書", 4, config.name, period_text(fin))
    labels = fin.result.chart.section_labels
    lines: list[Line] = []

    def section(key: str, rows: list[tuple[str, int]], total: int, sign: int) -> None:
        name = labels[key]
        lines.append(Line(f"    【{name}】"))
        lines.extend(Line(f"      {label}", v, 0) for label, v in rows)
        lines.append(Line(f"        {name}  計", sign * total, 1))

    section("revenue", fin.official_revenue_rows(), fin.revenue_total, 1)
    section("cost", fin.official_rows("cost"), fin.cost_total, -1)
    section("expense", fin.official_rows("expense"), fin.expense_total, -1)
    transfer_total = sum(v for _, v in fin.reserve_transfer_rows)
    lines.append(Line("    【正味財産】"))
    lines.append(Line("　　　当期正味財産増減額", fin.net_income, 1))
    lines.extend(Line(f"　　　{label}", v, 0) for label, v in fin.reserve_transfer_rows)
    lines.append(Line("        当期積立金振替額", -transfer_total, 1))
    lines.append(Line("　　　当期正味財産増減額－当期積立金振替額", fin.net_income - transfer_total, 1))
    ws.box(4, _lines(ws, lines, 4), 2, 4)
    return ws


def official_sheets(config: OrgConfig, fin: Financials) -> list[Sheet]:
    return [
        cover_sheet(config, fin),
        official_balance_sheet(config, fin),
        official_income_statement(config, fin),
    ]


# ---------------------------------------------------------------- ソフト様式


def software_balance_sheet(config: OrgConfig, fin: Financials) -> Sheet:
    left = [Line("  【流動資産】"), Line("    （現金・預金）")]
    left += [Line(f"      {a.name}", v) for a, v in fin.assets]
    left += [
        Line("        現金・預金 計", fin.asset_total),
        Line("          流動資産合計", fin.asset_total),
        Line("  【事業主貸】"),
        Line("      事業主  貸", 0),
    ]
    right = [Line("  【流動負債】")]
    right += [Line(f"      {a.name}", v) for a, v in fin.liabilities]
    right += [
        Line("          流動負債  計", fin.liability_total),
        Line("  【事業主借】"),
        Line("      事業主  借", 0),
        Line("負債の部合計", fin.liability_total, bold=True),
        Line("資　本　の　部", bold=True),
        Line("  【正味財産】"),
        Line(f"      {fin.result.chart.retained.name}", fin.retained_closing),
        Line(f"      {fin.result.chart.reserve.name} 合計", fin.reserve_total),
        # 第23期以降のソフト出力（科目性格の修正後）の表記に合わせる
        Line("      その他利益剰余金 合計", fin.reserve_total),
        Line("      利益剰余金 合計", fin.reserve_total),
        Line("          正味財産  計", fin.net_assets_subtotal),
        Line("  【所得】"),
        Line("      当期正味財産増減額", fin.net_income),
        Line("資本の部合計", fin.capital_total, bold=True),
    ]
    return _balance_sheet(
        config, fin, "負　債　・　資　本　の　部", left, right,
        Line("資産の部合計", fin.asset_total, bold=True),
        Line("負債・資本の部合計", fin.liabilities_and_capital_total, bold=True),
    )


def software_income_statement(config: OrgConfig, fin: Financials) -> Sheet:
    ws = Sheet("損益計算書")
    _widths(ws, {2: 53.0, 3: 20.0, 4: 20.0})
    _title(ws, "損　益　計　算　書", 4, config.name, period_text(fin))
    gross = fin.revenue_total - fin.cost_total
    operating = gross - fin.expense_total
    lines = [
        Line("《経常損益の部》"),
        Line("  ［営業損益の部］"),
        Line("    【売上高】"),
        *(Line(f"      {a.name}", v, 0) for a, v in fin.revenue),
        Line("        売上高  計", fin.revenue_total, 1),
        Line("    【売上原価】"),
        Line("        期首商品・製品棚卸高", 0, 0),
        *(Line(f"      {a.name}", v, 0) for a, v in fin.cost),
        Line("        当期商品仕入高", fin.cost_total, 0),
        Line("        合  計", fin.cost_total, 0),
        Line("        期末商品・製品棚卸高", 0, 0),
        Line("        売上原価  計", fin.cost_total, 1),
        Line("          売上総利益", gross, 1),
        Line("    【販売費一般管理費】"),
    ]
    if fin.expense:
        lines += [Line("      （その他経費）"), Line("          その他経費  計", fin.expense_total, 0)]
    lines += [
        Line("            販売費・一般管理費計", fin.expense_total, 0),
        Line("              営業利益", operating, 1),
        Line("  ［営業外損益の部］"),
        Line("    【営業外収益】"),
        Line("        営業外収益  計", 0, 0),
        Line("    【営業外費用】"),
        Line("        営業外費用  計", 0, 0),
        Line("          経常利益", operating, 1),
        Line("《引当金準備金》"),
        Line("    【繰戻額等】"),
        Line("        繰戻額等  計", 0, 0),
        Line("    【繰入額等】"),
        Line("        繰入額等  計", 0, 0),
        Line("          当期正味財産増減額", fin.net_income, 1),
    ]
    ws.box(4, _lines(ws, lines, 4), 2, 4)
    return ws


def software_expense_breakdown(config: OrgConfig, fin: Financials) -> Sheet:
    ws = Sheet("販売費及び一般管理費内訳書")
    _widths(ws, {2: 38.0, 3: 18.0, 4: 18.0})
    ws.merge(1, 2, 1, 4)
    ws.put(1, 2, "販売費及び一般管理費内訳書", bold=True, size=16, font=GOTHIC, align="center")
    ws.put(2, 4, TAX_NOTE, size=9, font=GOTHIC, align="right")
    ws.put(3, 2, config.name)
    ws.put(3, 4, period_text(fin), align="right")
    ws.underline(3, 2, 4)
    lines: list[Line] = []
    if fin.expense:
        lines.append(Line("（その他経費）"))
        lines += [Line(f"  {a.name}", v, 0) for a, v in fin.expense]
        lines.append(Line("      その他経費  計", fin.expense_total, 0))
    lines.append(Line("        販売費・一般管理費計", fin.expense_total, 0))
    ws.box(4, _lines(ws, lines, 4), 2, 3)
    return ws


def software_sheets(config: OrgConfig, fin: Financials) -> list[Sheet]:
    return [
        cover_sheet(config, fin),
        software_balance_sheet(config, fin),
        software_income_statement(config, fin),
        software_expense_breakdown(config, fin),
    ]


def software_values(fin: Financials) -> dict[str, int]:
    """ソフト様式の決算書の主な項目（ソフト出力の決算書と照合するためのもの）。"""
    values = {
        "資産の部合計": fin.asset_total,
        "流動負債計": fin.liability_total,
        "負債の部合計": fin.liability_total,
        fin.result.chart.retained.name: fin.retained_closing,
        f"{fin.result.chart.reserve.name}合計": fin.reserve_total,
        "正味財産計": fin.net_assets_subtotal,
        "当期正味財産増減額": fin.net_income,
        "資本の部合計": fin.capital_total,
        "負債・資本の部合計": fin.liabilities_and_capital_total,
        "売上高計": fin.revenue_total,
        "売上原価計": fin.cost_total,
        "販売費・一般管理費計": fin.expense_total,
    }
    for a, v in fin.assets + fin.liabilities + fin.revenue + fin.cost + fin.expense:
        values[a.software_name or a.name] = v
    return values


# ---------------------------------------------------------------- 書き出し


def save_xlsx(path: Path, sheets: list[Sheet]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin")
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in sheets:
        ws = wb.create_sheet(sheet.title)
        for (r, c), cell in sheet.cells.items():
            xc = ws.cell(row=r, column=c, value=cell.value)
            xc.font = Font(name=cell.font, size=cell.size, bold=cell.bold)
            if cell.align:
                xc.alignment = Alignment(horizontal=cell.align, vertical="center")
            if cell.number:
                xc.number_format = NUMBER_FORMAT
            if cell.borders:
                xc.border = Border(**{side: thin for side in cell.borders})
        for r1, c1, r2, c2 in sheet.merges:
            ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
        for col, width in sheet.widths.items():
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.page_setup.orientation = "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    wb.save(path)


def save_xls(path: Path, sheets: list[Sheet]) -> None:
    import xlwt

    styles: dict[tuple, xlwt.XFStyle] = {}

    def style(cell: Cell) -> xlwt.XFStyle:
        key = (cell.font, cell.size, cell.bold, cell.align, cell.number, frozenset(cell.borders))
        if key not in styles:
            st = xlwt.XFStyle()
            st.font.name = cell.font
            st.font.height = int(cell.size * 20)
            st.font.bold = cell.bold
            st.alignment.horz = {
                "left": xlwt.Alignment.HORZ_LEFT,
                "center": xlwt.Alignment.HORZ_CENTER,
                "right": xlwt.Alignment.HORZ_RIGHT,
            }.get(cell.align, xlwt.Alignment.HORZ_GENERAL)
            st.alignment.vert = xlwt.Alignment.VERT_CENTER
            for side in cell.borders:
                setattr(st.borders, side, xlwt.Borders.THIN)
            if cell.number:
                st.num_format_str = NUMBER_FORMAT
            styles[key] = st
        return styles[key]

    wb = xlwt.Workbook(encoding="utf-8")
    for sheet in sheets:
        ws = wb.add_sheet(sheet.title)
        merged = {}
        for r1, c1, r2, c2 in sheet.merges:
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    merged[(r, c)] = (r1, c1, r2, c2)
        for (r, c), cell in sheet.cells.items():
            if (r, c) in merged:
                r1, c1, r2, c2 = merged[(r, c)]
                if (r, c) == (r1, c1):
                    ws.write_merge(r1 - 1, r2 - 1, c1 - 1, c2 - 1, cell.value, style(cell))
                continue
            ws.write(r - 1, c - 1, cell.value, style(cell))
        for col, width in sheet.widths.items():
            ws.col(col - 1).width = int(width * 256)
        ws.portrait = True
        ws.paper_size_code = 9  # A4
        ws.fit_width_to_pages = 1
        ws.fit_height_to_pages = 0
    wb.save(str(path))


def write_official(path: Path, config: OrgConfig, fin: Financials) -> None:
    save_xlsx(path, official_sheets(config, fin))


def write_software(path: Path, config: OrgConfig, fin: Financials) -> None:
    save_xls(path, software_sheets(config, fin))
